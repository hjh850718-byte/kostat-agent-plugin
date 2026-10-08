#!/usr/bin/env python3
"""OOR 처리 결과 검증 모듈 (독립 실행 / 처리 후 호출).

기존 처리 로직은 수정하지 않는다. 원본 xlsx와 처리 결과 xlsx만 읽는다.

사용법:
    python oor_verify.py <이번주_원본.xlsx> <처리결과.xlsx> [--report <보고서.md>]

종료 코드:
    0  전 항목 PASS
    1  FAIL 1개 이상
    2  FAIL 없음 + SKIP(미확인) 있음   <- 0/1 외 코드. 정책 확정 필요

확정 전 항목 (CONFIG 참조, 해결하지 않고 위치만 표기):
    - Status 컬러 규칙: SKILL.md/references 의 색 기준과 충돌 가능. 기본값은 확정 루틴 Rule 5.
    - 기준 저장 경로 SAVE_DIR_REFERENCE: 미지정이면 항목 8 SKIP.
    - Status 열 STATUS_COL: 미지정이면 항목 5 SKIP.

가정 (샘플 파일 미확인 상태에서 정한 값):
    - 원본과 결과의 열 배치는 같다. A~R = 좌측 블록, AR = 마커, AS/AU = 참조 열.
    - 결과의 행 순서는 원본과 같다 (항목 9 행 단위 비교 조건).
    - 헤더는 1행.
"""
import argparse
import datetime as dt
import os
import re
import sys
from dataclasses import dataclass, field

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string as col_idx

CONFIG = {
    # 확정 전 (충돌 항목): 기본값은 확정 루틴 Rule 5 쪽
    "STATUS_COLOR_RULE": {
        "On-Time": None,  # 무색
        "Postpone process": "FFE0B2",
        "Bring process forward": "FFFF99",
        "Cancel": "FF0000",
    },
    "STATUS_COL": None,          # 확정 전: Status 열 문자. None이면 항목 5 SKIP
    "SAVE_DIR_REFERENCE": None,  # 확정 전: [경로] 기준 저장 경로. None이면 항목 8 SKIP
    "SPLIT_KEY_COL": None,       # 부분출하 분할 집계용 키 열. None이면 분할 설명 불가
    "HEADER_ROW": 1,
    "LEFT_LAST_COL": "R",        # 좌측 블록 마지막 열 (가정)
}

COL_PO_F, COL_PO_REF = "F", "AS"
COL_SHIP_H, COL_REQ_AU = "H", "AU"
COL_MARK = "AR"
COL_QTY = "M"
DATE_COLS = ["K", "L", "AX", "AY", "BB"]
HIDDEN_RANGES = [("A", "E"), ("I", "J"), ("N", "N"), ("P", "Q"), ("S", "AQ")]
FNAME_RE = re.compile(r"^OOR_Kostat_(\d{2})_(\d{2})_REV(\d{2})_updated\.xlsx$")
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


@dataclass
class Check:
    no: int
    name: str
    status: str  # PASS / FAIL / SKIP
    evidence: str
    details: list = field(default_factory=list)


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def fill_rgb(cell):
    f = cell.fill
    if f is None or f.fill_type != "solid" or not isinstance(f.fgColor.rgb, str):
        return None
    return f.fgColor.rgb[-6:].upper()


def last_data_row(path):
    usecols = list(range(col_idx(CONFIG["LEFT_LAST_COL"])))
    df = pd.read_excel(path, header=None, usecols=usecols, engine="openpyxl").dropna(how="all")
    return int(df.index.max()) + 1 if len(df) else CONFIG["HEADER_ROW"]


class Book:
    def __init__(self, path):
        self.path = path
        # keep_links=False: 외부 링크 갱신 방지. data_only=True: 셀 캐시 값 사용
        self.wb = load_workbook(path, keep_links=False, data_only=True)
        self.ws = self.wb.active
        self.last = last_data_row(path)
        self.rows = range(CONFIG["HEADER_ROW"] + 1, self.last + 1)
        self.n = self.last - CONFIG["HEADER_ROW"]

    def cell(self, letter, row):
        return self.ws.cell(row, col_idx(letter))

    def val(self, letter, row):
        return cell_text(self.cell(letter, row).value)


def check_count(o, r):
    delta = r.n - o.n
    if CONFIG["SPLIT_KEY_COL"]:
        k = CONFIG["SPLIT_KEY_COL"]
        ko = [o.val(k, x) for x in o.rows if o.val(k, x)]
        kr = [r.val(k, x) for x in r.rows if r.val(k, x)]
        split_extra = (len(kr) - len(set(kr))) - (len(ko) - len(set(ko)))
        status = "PASS" if delta == split_extra else "FAIL"
        ev = f"원본 {o.n}행 -> 결과 {r.n}행 (증가 {delta}). 분할 집계 {split_extra}행, 미설명 {delta - split_extra}행"
    else:
        split_extra = 0
        status = "PASS" if delta == 0 else "FAIL"
        ev = f"원본 {o.n}행 -> 결과 {r.n}행 (증가 {delta}). 분할 키 열 미지정으로 분할 설명 불가"
    return Check(1, "건수 대조", status, ev)


def check_alignment(r):
    bad = []
    for x in r.rows:
        f, a = r.val(COL_PO_F, x), r.val(COL_PO_REF, x)
        h, u = r.val(COL_SHIP_H, x), r.val(COL_REQ_AU, x)
        if f != a or h != u:
            bad.append(f"행 {x}: F={f} / AS={a}, H={h} / AU={u}")
    status = "PASS" if not bad else "FAIL"
    return Check(2, "정렬 정합 (F=AS, H=AU)", status, f"불일치 {len(bad)}행", bad)


def check_markers(r):
    h_only = [x for x in r.rows if r.val(COL_SHIP_H, x) and not r.val(COL_REQ_AU, x)]
    au_only = [x for x in r.rows if r.val(COL_REQ_AU, x) and not r.val(COL_SHIP_H, x)]
    marks = [x for x in r.rows if r.val(COL_MARK, x) == "≠"]
    det = [f"H only 행: {h_only}", f"AU only 행: {au_only}", f"AR '≠' 행: {marks}"]
    ok = not (h_only or au_only or marks)
    return Check(3, "H only / AU only / ≠ 마커", "PASS" if ok else "FAIL",
                 f"H only {len(h_only)} / AU only {len(au_only)} / ≠ {len(marks)}", det)


def check_dates(r):
    checked, bad = 0, []
    for letter in DATE_COLS:
        for x in r.rows:
            c = r.cell(letter, x)
            v = c.value
            if v is None or v == "":
                continue
            checked += 1
            if isinstance(v, dt.datetime):
                fmt_ok = c.number_format.lower() == "yyyy-mm-dd"
                if v.time() != dt.time(0) or not fmt_ok:
                    bad.append(f"{letter}{x}: datetime {v} / 표시형식 {c.number_format}")
            elif isinstance(v, dt.date):
                if c.number_format.lower() != "yyyy-mm-dd":
                    bad.append(f"{letter}{x}: date / 표시형식 {c.number_format}")
            elif not (isinstance(v, str) and DATE_RE.match(v.strip())):
                bad.append(f"{letter}{x}: '{v}'")
    return Check(4, "날짜 형식 (K,L,AX,AY,BB)", "PASS" if not bad else "FAIL",
                 f"검사 {checked}셀, 위반 {len(bad)}셀", bad[:20])


def check_status_colors(r):
    if not CONFIG["STATUS_COL"]:
        return Check(5, "Status 컬러 (확정 전)", "SKIP", "STATUS_COL 미지정 -> 검사 안 함")
    rule = CONFIG["STATUS_COLOR_RULE"]
    bad, unknown, checked = [], set(), 0
    for x in r.rows:
        status = r.val(CONFIG["STATUS_COL"], x)
        if not status:
            continue
        if status not in rule:
            unknown.add(status)
            continue
        checked += 1
        expected = rule[status]
        actual = fill_rgb(r.cell(CONFIG["STATUS_COL"], x))
        if (expected or None) != actual:
            bad.append(f"행 {x}: Status='{status}' 기대={expected or '무색'} 실제={actual or '무색'}")
    if unknown:
        bad.append(f"규칙에 없는 Status 값: {sorted(unknown)}")
    ev = f"검사 {checked}행, 불일치 {len(bad)}건 (규칙 확정 전, Status 셀 기준)"
    return Check(5, "Status 컬러 (확정 전)", "PASS" if not bad else "FAIL", ev, bad)


def check_hidden(r):
    hidden = set()
    for dim in r.ws.column_dimensions.values():
        if dim.hidden and dim.min and dim.max:
            hidden.update(range(dim.min, dim.max + 1))
    missing = []
    for start, end in HIDDEN_RANGES:
        for i in range(col_idx(start), col_idx(end) + 1):
            if i not in hidden:
                missing.append(i)
    from openpyxl.utils import get_column_letter
    names = [get_column_letter(i) for i in missing]
    return Check(6, "숨김 열 (A~E,I~J,N,P~Q,S~AQ)", "PASS" if not missing else "FAIL",
                 f"미숨김 열 {len(missing)}개", [f"미숨김: {names}"] if missing else [])


def check_filename(res_path):
    base = os.path.basename(res_path)
    m = FNAME_RE.match(base)
    if not m:
        return Check(7, "파일명 형식 / REV", "FAIL", f"형식 불일치: {base}")
    rev = int(m.group(3))
    folder = os.path.dirname(os.path.abspath(res_path))
    prev = []
    for name in os.listdir(folder):
        pm = FNAME_RE.match(name)
        if pm and name != base:
            prev.append(int(pm.group(3)))
    if not prev:
        return Check(7, "파일명 형식 / REV", "SKIP", f"형식 OK, 같은 폴더에 기존 REV 파일 없음 (현재 REV{rev:02d})")
    expected = max(prev) + 1
    status = "PASS" if rev == expected else "FAIL"
    return Check(7, "파일명 형식 / REV", status,
                 f"형식 OK, REV{rev:02d} (기대 REV{expected:02d} = 기존 최대 REV{max(prev):02d} + 1)")


def check_save_dir(res_path):
    ref = CONFIG["SAVE_DIR_REFERENCE"]
    if not ref:
        return Check(8, "저장 경로 (확정 전)", "SKIP", "SAVE_DIR_REFERENCE 미지정 [경로] -> 검사 안 함")
    actual = os.path.normcase(os.path.abspath(os.path.dirname(res_path)))
    expected = os.path.normcase(os.path.abspath(ref))
    return Check(8, "저장 경로 (확정 전)", "PASS" if actual == expected else "FAIL",
                 f"실제 {actual} / 기준 {expected}")


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def check_qty(o, r):
    sum_o = sum(n for n in (_num(o.val(COL_QTY, x)) for x in o.rows) if n is not None)
    sum_r = sum(n for n in (_num(r.val(COL_QTY, x)) for x in r.rows) if n is not None)
    det = [f"합계: 원본 {sum_o:g} / 결과 {sum_r:g}"]
    ok = abs(sum_o - sum_r) < 1e-9
    if o.n == r.n:
        mism = [f"행 {a}: 원본 {o.val(COL_QTY, a)} / 결과 {r.val(COL_QTY, b)}"
                for a, b in zip(o.rows, r.rows) if o.val(COL_QTY, a) != r.val(COL_QTY, b)]
        det += mism
        ok = ok and not mism
    else:
        det.append("행 수가 달라 행 단위 비교 생략 (합계만 비교)")
    return Check(9, "Qty(M열) 원본 동일", "PASS" if ok else "FAIL",
                 f"합계 원본 {sum_o:g} / 결과 {sum_r:g}", det)


def run_checks(orig_path, res_path):
    o, r = Book(orig_path), Book(res_path)
    return [
        check_count(o, r),
        check_alignment(r),
        check_markers(r),
        check_dates(r),
        check_status_colors(r),
        check_hidden(r),
        check_filename(res_path),
        check_save_dir(res_path),
        check_qty(o, r),
    ]


def print_table(checks):
    print(f"{'#':>2} | {'항목':<30} | {'결과':<4} | 근거")
    print("-" * 110)
    for c in checks:
        print(f"{c.no:>2} | {c.name:<30} | {c.status:<4} | {c.evidence}")
        for line in c.details[:10]:
            print(f"   |   - {line}")


def write_report(path, orig_path, res_path, checks):
    counts = {s: sum(1 for c in checks if c.status == s) for s in ("PASS", "FAIL", "SKIP")}
    lines = [
        "# OOR 처리 결과 검증 리포트",
        "",
        f"- 생성 시각: {dt.datetime.now():%Y-%m-%d %H:%M:%S}",
        f"- 원본: `{os.path.basename(orig_path)}`",
        f"- 결과: `{os.path.basename(res_path)}`",
        f"- 요약: PASS {counts['PASS']} / FAIL {counts['FAIL']} / SKIP {counts['SKIP']}",
        "- **확정 전 항목**: Status 컬러(5), 저장 경로(8). 기본값은 확정 루틴 Rule 5 쪽이며 충돌 항목으로 미확정.",
        "",
        "| # | 항목 | 결과 | 근거 |",
        "|---|------|------|------|",
    ]
    for c in checks:
        lines.append(f"| {c.no} | {c.name} | {c.status} | {c.evidence} |")
    for c in checks:
        if c.details:
            lines += ["", f"## {c.no}. {c.name}", ""]
            lines += [f"- {d}" for d in c.details]
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")


def main(argv=None):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except AttributeError:
        pass
    ap = argparse.ArgumentParser(description="OOR 처리 결과 검증")
    ap.add_argument("orig", help="이번 주 원본 xlsx")
    ap.add_argument("result", help="처리 결과 xlsx")
    ap.add_argument("--report", help="리포트 저장 경로 (기본: 결과 파일과 같은 폴더의 oor_verify_report.md)")
    args = ap.parse_args(argv)

    checks = run_checks(args.orig, args.result)
    print_table(checks)

    report = args.report or os.path.join(os.path.dirname(os.path.abspath(args.result)), "oor_verify_report.md")
    write_report(report, args.orig, args.result, checks)
    print(f"\n리포트 저장: {report}")

    if any(c.status == "FAIL" for c in checks):
        return 1
    if any(c.status == "SKIP" for c in checks):
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
